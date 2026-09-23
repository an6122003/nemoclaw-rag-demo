#!/usr/bin/env python3
"""Generate the Hands-on 3 sales workbook: aurora_sales_2024_2025.xlsx.

Aurora Grid Systems is the same fictional company as the Hands-on 2 corpus. This
workbook is its Vietnam sales ledger for 2024-2025, one row per month, region,
product and sales channel.

The numbers are synthetic but deliberately shaped, the same way the Hands-on 2
corpus hides retrieval traps. A good analyst (human or agent) has to notice:

  * The biggest region is not the fastest-growing one. TP. Ho Chi Minh has the
    most revenue; Da Nang grows fastest (new solar-plus-storage projects from
    Q2 2025); Ha Noi shrinks, with a visible Q3 2025 dip.
  * "Best-selling" depends on the metric. The residential Aurora Home 10 battery
    sells the most units; the AX-600 cabinet earns the most revenue.
  * AX-400 sales fall after April 2025 - the month safety bulletin SB-2025-04
    (AX-400 thermal derate) was issued in the Hands-on 2 corpus.
  * Q4 is the strongest quarter every year; February (Tet) is the weakest month.

Deterministic: the same seed always produces the same workbook, so the answer
key in README.md stays valid. Requires pandas and openpyxl.

    python3 make_dataset.py              # writes aurora_sales_2024_2025.xlsx here
    python3 make_dataset.py --check      # print the headline facts only
"""

from __future__ import annotations

import argparse
import random
from datetime import date
from pathlib import Path

import pandas as pd

OUT = Path(__file__).resolve().parent / "aurora_sales_2024_2025.xlsx"
SEED = 2025

# Average monthly revenue per region in 2024 (million VND), and the shape of
# 2025 relative to the same month of 2024.
REGIONS = {
    "TP. Hồ Chí Minh": 41_000,
    "Hà Nội": 30_500,
    "Hải Phòng": 13_200,
    "Đà Nẵng": 9_400,
    "Cần Thơ": 5_600,
}


def region_2025_factor(region: str, month: int) -> float:
    """Multiplier for a 2025 month versus the same month in 2024."""
    if region == "TP. Hồ Chí Minh":
        return 1.12
    if region == "Hà Nội":
        # Flat-to-soft year with a clear Q3 dip (a large EPC contract slipped).
        return 0.80 if month in (7, 8, 9) else 1.00
    if region == "Hải Phòng":
        return 1.17
    if region == "Đà Nẵng":
        # New solar-plus-storage projects ramp up from Q2 2025.
        return 1.18 if month <= 3 else 1.80
    if region == "Cần Thơ":
        return 1.20 + 0.025 * month  # steady acceleration through the year
    raise KeyError(region)


# Month-of-year seasonality (Tet in February, year-end project completions).
SEASON = {1: 0.82, 2: 0.66, 3: 0.95, 4: 1.00, 5: 1.02, 6: 1.06,
          7: 0.96, 8: 0.95, 9: 1.00, 10: 1.10, 11: 1.20, 12: 1.28}
_mean = sum(SEASON.values()) / 12
SEASON = {m: v / _mean for m, v in SEASON.items()}

# name: (product group, unit price in million VND, gross margin, revenue share 2024)
PRODUCTS = {
    "Tủ pin AX-600": ("Thương mại", 3_950, 0.22, 0.34),
    "Tủ pin AX-400": ("Thương mại", 2_750, 0.18, 0.26),
    "Bộ chuyển đổi Helios H3": ("Thương mại", 1_480, 0.15, 0.20),
    "Pin gia đình Aurora Home 10": ("Dân dụng", 85, 0.27, 0.12),
    "Dịch vụ bảo trì": ("Dịch vụ", 120, 0.46, 0.08),
}


def product_share(product: str, year: int, month: int) -> float:
    """Revenue share of a product within a region-month, before renormalising."""
    base = PRODUCTS[product][3]
    if year == 2024:
        return base
    if product == "Tủ pin AX-600":
        return base * 1.25
    if product == "Tủ pin AX-400":
        # SB-2025-04 (AX-400 thermal derate) was issued in April 2025.
        return base * (0.95 if month < 4 else 0.52)
    if product == "Pin gia đình Aurora Home 10":
        return base * 1.30
    if product == "Dịch vụ bảo trì":
        return base * 1.15
    return base


# channel mix per product group
CHANNELS = {
    "Thương mại": {"Bán trực tiếp": 0.50, "Đối tác EPC": 0.42, "Đại lý": 0.08},
    "Dân dụng": {"Bán trực tiếp": 0.12, "Đối tác EPC": 0.08, "Đại lý": 0.80},
    "Dịch vụ": {"Bán trực tiếp": 0.85, "Đối tác EPC": 0.15},
}

COLUMNS = [
    "Ngày", "Năm", "Quý", "Tháng", "Khu vực", "Sản phẩm", "Nhóm sản phẩm",
    "Kênh bán", "Số lượng", "Doanh thu (triệu VND)", "Giá vốn (triệu VND)",
    "Lợi nhuận gộp (triệu VND)",
]


def build() -> pd.DataFrame:
    rng = random.Random(SEED)
    rows = []
    # Error-diffusion rounding: each (region, product, channel) carries its
    # fractional unit forward, so integer units still add up to the designed
    # totals. Independent random rounding made small regions far too noisy -
    # one AX-600 more or less swung Can Tho's growth by 20 points.
    carry: dict[tuple[str, str, str], float] = {}
    for year in (2024, 2025):
        for month in range(1, 13):
            for region, base in REGIONS.items():
                total = base * SEASON[month]
                if year == 2025:
                    total *= region_2025_factor(region, month)
                total *= rng.uniform(0.94, 1.06)

                shares = {p: product_share(p, year, month) for p in PRODUCTS}
                norm = sum(shares.values())
                for product, (group, price, margin, _) in PRODUCTS.items():
                    product_rev = total * shares[product] / norm
                    for channel, cshare in CHANNELS[group].items():
                        expected = product_rev * cshare
                        key = (region, product, channel)
                        # Start each carry at a random phase. Starting at zero
                        # delays a rarely sold cabinet's first unit, which moved
                        # 2024 revenue into 2025 and inflated growth.
                        if key not in carry:
                            carry[key] = rng.random()
                        carry[key] += expected / price
                        units = int(carry[key])
                        carry[key] -= units
                        if units == 0:
                            continue
                        revenue = units * price * rng.uniform(0.96, 1.00)  # discounts
                        cost = revenue * (1 - margin * rng.uniform(0.92, 1.08))
                        rows.append({
                            "Ngày": date(year, month, 1),
                            "Năm": year,
                            "Quý": f"Q{(month - 1) // 3 + 1}",
                            "Tháng": month,
                            "Khu vực": region,
                            "Sản phẩm": product,
                            "Nhóm sản phẩm": group,
                            "Kênh bán": channel,
                            "Số lượng": units,
                            "Doanh thu (triệu VND)": round(revenue, 1),
                            "Giá vốn (triệu VND)": round(cost, 1),
                            "Lợi nhuận gộp (triệu VND)": round(revenue - cost, 1),
                        })
    df = pd.DataFrame(rows, columns=COLUMNS)
    df["Ngày"] = pd.to_datetime(df["Ngày"])
    return df


def headline(df: pd.DataFrame) -> list[str]:
    rev = "Doanh thu (triệu VND)"
    out = []
    by = df.pivot_table(index="Khu vực", columns="Năm", values=rev, aggfunc="sum")
    by["growth"] = by[2025] / by[2024] - 1
    out.append(f"rows: {len(df)}")
    out.append(f"total revenue 2024: {df[df['Năm'] == 2024][rev].sum() / 1000:,.1f} bn VND, "
               f"2025: {df[df['Năm'] == 2025][rev].sum() / 1000:,.1f} bn VND")
    out.append("largest region 2025: " + by[2025].idxmax())
    out.append("fastest growth: " + by["growth"].idxmax()
               + f" ({by['growth'].max():+.1%})")
    out.append("slowest growth: " + by["growth"].idxmin()
               + f" ({by['growth'].min():+.1%})")
    for region, row in by.sort_values(2025, ascending=False).iterrows():
        out.append(f"  {region:18s} 2024 {row[2024] / 1000:7.1f} bn  "
                   f"2025 {row[2025] / 1000:7.1f} bn  {row['growth']:+.1%}")
    prod = df.groupby("Sản phẩm")[["Số lượng", rev]].sum()
    out.append("most units: " + prod["Số lượng"].idxmax())
    out.append("most revenue: " + prod[rev].idxmax())
    q = df.groupby(["Năm", "Quý"])[rev].sum()
    out.append("strongest quarter per year: "
               + ", ".join(f"{y} {q[y].idxmax()}" for y in (2024, 2025)))
    ax4 = df[df["Sản phẩm"] == "Tủ pin AX-400"].groupby(["Năm", "Tháng"])[rev].sum()
    out.append(f"AX-400 2025 Jan-Mar avg {ax4[2025].loc[1:3].mean() / 1000:.1f} bn/month, "
               f"Apr-Dec avg {ax4[2025].loc[4:12].mean() / 1000:.1f} bn/month")
    return out


def write_workbook(df: pd.DataFrame, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    about = pd.DataFrame({
        "Cột": COLUMNS,
        "Ý nghĩa": [
            "Tháng ghi nhận doanh số (ngày đầu tháng)",
            "Năm", "Quý", "Tháng (1-12)",
            "Khu vực bán hàng", "Tên sản phẩm hoặc dịch vụ",
            "Thương mại / Dân dụng / Dịch vụ",
            "Bán trực tiếp / Đối tác EPC / Đại lý",
            "Số đơn vị bán ra (tủ, bộ, pin, hoặc hợp đồng bảo trì)",
            "Doanh thu sau chiết khấu, triệu đồng",
            "Giá vốn hàng bán, triệu đồng",
            "Doanh thu trừ giá vốn, triệu đồng",
        ],
    })
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        df.to_excel(xw, sheet_name="Dữ liệu bán hàng", index=False)
        about.to_excel(xw, sheet_name="Giới thiệu", index=False)
        ws = xw.sheets["Dữ liệu bán hàng"]
        head = PatternFill("solid", fgColor="0B6E4F")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        widths = [12, 7, 6, 7, 18, 28, 14, 15, 10, 16, 16, 18]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
            row[0].number_format = "yyyy-mm"
        for col in (10, 11, 12):
            for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
                row[0].number_format = "#,##0.0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        about_ws = xw.sheets["Giới thiệu"]
        about_ws.column_dimensions["A"].width = 28
        about_ws.column_dimensions["B"].width = 60
        for cell in about_ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    ap.add_argument("--out", default=str(OUT))
    ap.add_argument("--check", action="store_true", help="print headline facts, write nothing")
    args = ap.parse_args()
    df = build()
    for line in headline(df):
        print(line)
    if not args.check:
        write_workbook(df, Path(args.out))
        print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
