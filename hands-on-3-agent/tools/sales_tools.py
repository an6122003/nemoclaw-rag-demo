#!/usr/bin/env python3
"""Sales-analyst tools for the Hands-on 3 agent: pandas, matplotlib, openpyxl.

The agent (OpenClaw inside the NemoClaw sandbox) decides WHICH tool to call and
with WHAT arguments. This module does the actual work, and it never asks the
language model to do arithmetic: every number the agent needs comes back
pre-computed and pre-formatted. That split - the model chooses, the tool
computes - is the main design lesson of the lab.

Four tools, described to the model by TOOL_SPECS (OpenAI function-calling
schema):

    get_dataset_info     read the workbook: sheets, columns, types, time range
    analyze_sales        pandas group-by, per-year totals, growth, share, rank
    create_chart         matplotlib chart (bar / line / pie) saved as PNG
    export_excel_report  openpyxl .xlsx report: summary, table + native chart,
                         chart image, underlying rows

Works with any tabular .xlsx or .csv, not only the bundled workbook: column
names are resolved loosely (case, accents and common synonyms such as
"doanh số" -> "Doanh thu (triệu VND)"), and a wrong name returns the list of
real columns so the agent can correct itself.

Two ways in:

  * imported by the workshop app, which executes the agent's tool calls, and
  * a command line, used by the OpenClaw skill inside the sandbox:

        python3 sales_tools.py get_dataset_info --data sales.xlsx
        python3 sales_tools.py analyze_sales --data sales.xlsx \\
            --args '{"group_by": "Khu vực", "metric": "doanh thu"}'
        python3 sales_tools.py specs            # print the tool schemas

Requires pandas, matplotlib and openpyxl (see ../requirements.txt).
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import threading
import time
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pandas as pd

HERE = Path(__file__).resolve().parent
DEFAULT_DATA = HERE.parent / "data" / "aurora_sales_2024_2025.xlsx"
MAX_MODEL_ROWS = 15  # rows returned to the model; the UI gets all of them

# --------------------------------------------------------------------------
# Tool schemas shown to the model
# --------------------------------------------------------------------------
_GROUP_BY = {
    "type": "string",
    "description": (
        "Column to group by, exactly as listed by get_dataset_info. For the bundled "
        "workbook: 'Khu vực' (region/vùng), 'Sản phẩm' (product), 'Nhóm sản phẩm' "
        "(product group), 'Kênh bán' (sales channel), 'Tháng' (month), 'Quý' (quarter), "
        "'Năm' (year)."
    ),
}
_METRIC = {
    "type": "string",
    "description": (
        "What to measure. A numeric column, e.g. 'Doanh thu (triệu VND)' for revenue / "
        "doanh số / doanh thu, 'Lợi nhuận gộp (triệu VND)' for profit / lợi nhuận, "
        "'Số lượng' for units sold / sản lượng. Use 'margin' for gross margin % "
        "(biên lợi nhuận)."
    ),
}
_FILTERS = {
    "type": "object",
    "description": (
        "Optional filters, column -> value, e.g. {\"Khu vực\": \"Đà Nẵng\"} or "
        "{\"Năm\": 2025}. Omit to use all rows."
    ),
    "additionalProperties": {"type": ["string", "number"]},
}

TOOL_SPECS: list[dict] = [
    {
        "type": "function",
        "function": {
            "name": "get_dataset_info",
            "description": (
                "Open the sales Excel file and describe it: sheets, columns with their "
                "types and example values, number of rows, and the time range. Call this "
                "first, before any analysis, to learn the real column names."
            ),
            "parameters": {"type": "object", "properties": {}, "required": []},
        },
    },
    {
        "type": "function",
        "function": {
            "name": "analyze_sales",
            "description": (
                "Python + pandas analysis. Groups the rows by one column, totals a metric "
                "for each year, and computes year-over-year growth, share of total and "
                "ranking. Returns a table plus pre-computed highlights (largest, fastest "
                "growth, slowest growth, total). Use it for every question about totals, "
                "rankings, growth, trends or comparisons. Quote its numbers exactly."
            ),
            "parameters": {
                "type": "object",
                "properties": {"group_by": _GROUP_BY, "metric": _METRIC, "filters": _FILTERS},
                "required": ["group_by", "metric"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "create_chart",
            "description": (
                "Draw a chart with matplotlib and save it as a PNG image. Use chart_type "
                "'bar' to compare groups year by year, 'line' for a trend over months or "
                "quarters, 'pie' for the share of the latest year."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "group_by": _GROUP_BY,
                    "metric": _METRIC,
                    "chart_type": {"type": "string", "enum": ["bar", "line", "pie"]},
                    "title": {"type": "string", "description": "Chart title, in the user's language."},
                    "filters": _FILTERS,
                },
                "required": ["group_by", "metric", "chart_type"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "export_excel_report",
            "description": (
                "Create a formatted Excel (.xlsx) statistics report: a summary sheet with "
                "your key insights, the analysis table with a native Excel chart, the "
                "chart image, and the underlying rows. Call it after analyze_sales."
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Report title, in the user's language."},
                    "group_by": _GROUP_BY,
                    "metric": _METRIC,
                    "filters": _FILTERS,
                    "insights": {
                        "type": "array",
                        "items": {"type": "string"},
                        "description": "3-5 short key findings for the summary sheet, in the user's language.",
                    },
                },
                "required": ["title", "group_by", "metric", "insights"],
            },
        },
    },
]

TOOL_NAMES = [t["function"]["name"] for t in TOOL_SPECS]


# --------------------------------------------------------------------------
# Context and errors
# --------------------------------------------------------------------------
class ToolError(Exception):
    """A problem the agent can fix by calling again with better arguments."""

    def __init__(self, message: str, **extra: Any):
        super().__init__(message)
        self.extra = extra


@dataclass
class ToolContext:
    """Per-run state. Created by the caller, never seen by the model."""

    data_path: Path = DEFAULT_DATA
    out_dir: Path = Path("outputs")
    lang: str = "vi"
    counter: int = 0
    charts: list[dict] = field(default_factory=list)
    files: list[dict] = field(default_factory=list)

    def next_id(self, prefix: str) -> str:
        self.counter += 1
        return f"{prefix}{self.counter}"


# --------------------------------------------------------------------------
# Text helpers
# --------------------------------------------------------------------------
def norm(text: Any) -> str:
    """Lowercase, strip accents (đ -> d) and punctuation: 'Khu vực' -> 'khu vuc'."""
    s = unicodedata.normalize("NFD", str(text).lower().replace("đ", "d"))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


# concept -> synonyms (normalized). Used to resolve loose names such as
# "doanh số" or "region" to the workbook's real column.
CONCEPTS: dict[str, list[str]] = {
    "region": ["khu vuc", "vung", "region", "area", "tinh thanh", "tinh", "chi nhanh", "mien"],
    "product": ["san pham", "product", "mat hang", "item"],
    "group": ["nhom san pham", "nhom hang", "product group", "category", "nganh hang"],
    "channel": ["kenh ban", "kenh ban hang", "kenh", "channel"],
    "month": ["thang", "month"],
    "quarter": ["quy", "quarter"],
    "year": ["nam", "year"],
    "date": ["ngay", "date", "thoi gian"],
    "revenue": ["doanh thu", "doanh so", "revenue", "sales", "turnover", "ban hang"],
    "profit": ["loi nhuan gop", "loi nhuan", "lai gop", "profit", "gross profit", "margin value"],
    "units": ["so luong", "san luong", "units", "quantity", "qty", "volume", "so don vi"],
    "cost": ["gia von", "chi phi", "cost", "cogs"],
}
MARGIN_WORDS = {"margin", "bien loi nhuan", "ty suat loi nhuan", "gross margin", "bien lai",
                "ty le loi nhuan", "margin pct", "margin percent"}

PLACE_ALIASES = {
    "hcm": "ho chi minh", "tphcm": "ho chi minh", "tp hcm": "ho chi minh",
    "sai gon": "ho chi minh", "saigon": "ho chi minh", "hn": "ha noi",
    "hp": "hai phong", "dn": "da nang", "ct": "can tho",
}

MONTHS_EN = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

T = {
    "vi": {
        "growth": "tăng trưởng", "share": "tỷ trọng", "rank": "hạng", "total": "Tổng",
        "value": "giá trị", "largest": "Lớn nhất", "fastest": "Tăng trưởng nhanh nhất",
        "slowest": "Tăng trưởng chậm nhất", "smallest": "Nhỏ nhất", "vs": "so với",
        "least_decline": "Giảm ít nhất", "most_decline": "Giảm mạnh nhất",
        "change": "chênh lệch", "biggest_gain": "Tăng nhiều nhất về giá trị",
        "biggest_loss": "Giảm nhiều nhất về giá trị", "units_word": "đơn vị",
        "best_change": "Cải thiện nhiều nhất", "worst_change": "Kém cải thiện nhất",
        "of_total": "tổng", "month": "Tháng", "no_growth": "không đủ dữ liệu để tính tăng trưởng",
        "source": "Nguồn", "made_by": "Tạo tự động bởi Sales Analyst Agent · NemoClaw trên DGX Spark",
        "summary": "Tóm tắt", "analysis": "Phân tích", "chart": "Biểu đồ", "data": "Dữ liệu",
        "insights": "Nhận định chính", "kpis": "Chỉ số chính", "filters": "Bộ lọc",
        "all_rows": "toàn bộ dữ liệu", "generated": "Thời điểm tạo", "dataset": "Tệp dữ liệu",
        "group_by": "Nhóm theo", "metric": "Chỉ tiêu", "margin": "Biên lợi nhuận gộp",
        "pp": "điểm %", "rows_used": "Số dòng dữ liệu", "truncated": "Chỉ hiển thị {n} dòng đầu",
        "billion": "tỷ đồng", "million": "triệu đồng",
    },
    "en": {
        "growth": "growth", "share": "share", "rank": "rank", "total": "Total",
        "value": "value", "largest": "Largest", "fastest": "Fastest growth",
        "slowest": "Slowest growth", "smallest": "Smallest", "vs": "vs",
        "least_decline": "Smallest decline", "most_decline": "Largest decline",
        "change": "change", "biggest_gain": "Largest increase in value",
        "biggest_loss": "Largest decrease in value", "units_word": "units",
        "best_change": "Most improved", "worst_change": "Least improved",
        "of_total": "of total", "month": "Month", "no_growth": "not enough data to compute growth",
        "source": "Source", "made_by": "Generated by the Sales Analyst Agent · NemoClaw on DGX Spark",
        "summary": "Summary", "analysis": "Analysis", "chart": "Chart", "data": "Data",
        "insights": "Key insights", "kpis": "Key figures", "filters": "Filters",
        "all_rows": "all rows", "generated": "Generated", "dataset": "Dataset",
        "group_by": "Grouped by", "metric": "Metric", "margin": "Gross margin",
        "pp": "pp", "rows_used": "Rows used", "truncated": "Showing the first {n} rows",
        "billion": "bn VND", "million": "M VND",
    },
}


def tr(ctx: ToolContext, key: str) -> str:
    return T.get(ctx.lang, T["en"]).get(key, T["en"][key])


# --------------------------------------------------------------------------
# Number formatting (the model copies these strings verbatim)
# --------------------------------------------------------------------------
def _group_digits(value: float, decimals: int, lang: str) -> str:
    s = f"{abs(value):,.{decimals}f}"
    if lang == "vi":  # 1,234.5 -> 1.234,5
        s = s.replace(",", "\x00").replace(".", ",").replace("\x00", ".")
    return ("-" if value < 0 else "") + s


def fmt_pct(ratio: float | None, lang: str, signed: bool = True) -> str:
    if ratio is None or (isinstance(ratio, float) and math.isnan(ratio)):
        return "—"
    s = _group_digits(ratio * 100, 1, lang)
    if signed and ratio > 0:
        s = "+" + s
    return s + "%"


def detect_unit(column: str) -> str:
    """Return 'million_vnd', 'vnd', 'billion_vnd' or '' from a column name."""
    n = norm(column)
    if "vnd" in n or "dong" in n or "vnđ" in column.lower():
        if "trieu" in n:
            return "million_vnd"
        if "ty" in n.split():
            return "billion_vnd"
        return "vnd"
    return ""


def fmt_value(value: float, unit: str, lang: str) -> str:
    """Human-readable money or count, e.g. 184485.2 (million VND) -> '184,5 tỷ đồng'."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return "—"
    t = T.get(lang, T["en"])
    if unit == "million_vnd":
        if abs(value) >= 1000:
            return f"{_group_digits(value / 1000, 1, lang)} {t['billion']}"
        return f"{_group_digits(value, 0, lang)} {t['million']}"
    if unit == "billion_vnd":
        return f"{_group_digits(value, 1, lang)} {t['billion']}"
    if unit == "vnd":
        if abs(value) >= 1e9:
            return f"{_group_digits(value / 1e9, 1, lang)} {t['billion']}"
        if abs(value) >= 1e6:
            return f"{_group_digits(value / 1e6, 0, lang)} {t['million']}"
        return _group_digits(value, 0, lang) + " VND"
    decimals = 0 if float(value).is_integer() or abs(value) >= 100 else 1
    return _group_digits(value, decimals, lang)


# --------------------------------------------------------------------------
# Dataset loading and column resolution
# --------------------------------------------------------------------------
_cache: dict[tuple[str, float], tuple[pd.DataFrame, dict]] = {}
_cache_lock = threading.Lock()


def load_dataset(path: Path) -> tuple[pd.DataFrame, dict]:
    """Load the largest table in an .xlsx/.csv file. Cached by path and mtime."""
    path = Path(path)
    if not path.exists():
        raise ToolError(f"data file not found: {path.name}")
    key = (str(path), path.stat().st_mtime)
    with _cache_lock:
        if key in _cache:
            df, meta = _cache[key]
            return df.copy(), dict(meta)

    suffix = path.suffix.lower()
    sheets: list[str] = []
    if suffix == ".csv":
        df = pd.read_csv(path)
        sheet = path.stem
        sheets = [sheet]
    elif suffix in (".xlsx", ".xlsm"):
        frames = pd.read_excel(path, sheet_name=None)
        sheets = list(frames)
        usable = {k: v for k, v in frames.items() if v.shape[1] >= 2 and len(v) > 0}
        if not usable:
            raise ToolError("the workbook has no table with at least two columns")
        sheet = max(usable, key=lambda k: len(usable[k]))
        df = usable[sheet]
    else:
        raise ToolError(f"unsupported file type {suffix}; use .xlsx or .csv")

    df = df.dropna(how="all").dropna(axis=1, how="all")
    df.columns = [str(c).strip() for c in df.columns]
    meta = {"file": path.name, "sheet": sheet, "sheets": sheets}
    meta.update(_detect_time(df))
    with _cache_lock:
        _cache[key] = (df, meta)
    return df.copy(), dict(meta)


def _concept_of(column: str) -> str | None:
    n = norm(column)
    best, best_len = None, 0
    for concept, words in CONCEPTS.items():
        for w in words:
            if n == w or n.startswith(w + " ") or f" {w} " in f" {n} ":
                if len(w) > best_len:
                    best, best_len = concept, len(w)
    return best


def _detect_time(df: pd.DataFrame) -> dict:
    """Find a year column (or derive one from a date column)."""
    year_col = date_col = None
    for c in df.columns:
        concept = _concept_of(c)
        if concept == "year" and pd.api.types.is_numeric_dtype(df[c]):
            vals = df[c].dropna()
            if len(vals) and vals.between(1990, 2100).all():
                year_col = c
        if concept == "date" or pd.api.types.is_datetime64_any_dtype(df[c]):
            if pd.api.types.is_datetime64_any_dtype(df[c]):
                date_col = date_col or c
    return {"year_col": year_col, "date_col": date_col}


def resolve_column(df: pd.DataFrame, name: str, numeric: bool | None = None) -> str:
    """Map a loose name ('doanh số', 'region', 'khu vuc') to a real column."""
    if name in df.columns:
        return name
    q = norm(name)
    cols = list(df.columns)

    def ok(c: str) -> bool:
        if numeric is None:
            return True
        return pd.api.types.is_numeric_dtype(df[c]) == numeric

    exact = [c for c in cols if norm(c) == q and ok(c)]
    if exact:
        return exact[0]
    # 'doanh thu' -> 'Doanh thu (triệu VND)'
    prefix = [c for c in cols if ok(c) and (norm(c).startswith(q + " ") or q.startswith(norm(c) + " "))]
    if prefix:
        return min(prefix, key=lambda c: len(norm(c)))
    # synonyms: 'doanh số' and 'revenue' both mean the revenue concept
    wanted = None
    for concept, words in CONCEPTS.items():
        if q in words or any(q == w or q.startswith(w + " ") for w in words):
            wanted = concept
            break
    if wanted:
        matches = [c for c in cols if ok(c) and _concept_of(c) == wanted]
        if matches:
            return min(matches, key=lambda c: len(norm(c)))
    contains = [c for c in cols if ok(c) and q and (q in norm(c))]
    if contains:
        return min(contains, key=lambda c: len(norm(c)))
    kind = "numeric " if numeric else ""
    raise ToolError(
        f"no {kind}column matches '{name}'",
        available_columns=[c for c in cols if ok(c)] or cols,
        hint="Call again with one of available_columns, spelled exactly.",
    )


def _is_margin(metric: str) -> bool:
    return norm(metric) in MARGIN_WORDS


def _match_value(series: pd.Series, value: Any) -> pd.Series:
    """Exact-ish filter match: numbers numerically, text ignoring case/accents."""
    if isinstance(value, (list, tuple)):
        mask = pd.Series(False, index=series.index)
        for v in value:
            mask |= _match_value(series, v)
        return mask
    if pd.api.types.is_numeric_dtype(series):
        try:
            return series == float(value)
        except (TypeError, ValueError):
            return pd.Series(False, index=series.index)
    q = norm(value)
    q = PLACE_ALIASES.get(q, q)
    normed = series.astype(str).map(norm)
    exact = normed == q
    if exact.any():
        return exact
    return normed.str.contains(re.escape(q), regex=True) if q else exact


def apply_filters(df: pd.DataFrame, filters: dict | None) -> tuple[pd.DataFrame, dict]:
    applied: dict[str, Any] = {}
    for key, value in (filters or {}).items():
        if value is None or value == "":
            continue
        col = resolve_column(df, str(key))
        mask = _match_value(df[col], value)
        if not mask.any():
            sample = sorted({str(v) for v in df[col].dropna().unique()})[:20]
            raise ToolError(f"no rows where {col} = {value!r}", available_values=sample,
                            hint="Use one of available_values.")
        df = df[mask]
        applied[col] = value
    return df, applied


# --------------------------------------------------------------------------
# Analysis core (shared by analyze / chart / report)
# --------------------------------------------------------------------------
@dataclass
class Analysis:
    group_col: str
    metric_col: str            # a real column, or "margin"
    metric_label: str
    unit: str                  # million_vnd | vnd | pct | ''
    years: list                # [] when there is no year dimension
    table: pd.DataFrame        # index = group label; columns = years (or 'value'), growth, share
    filters: dict
    rows_used: int
    time_like: bool
    frame: pd.DataFrame        # filtered raw rows


def _time_like(col: str) -> bool:
    return _concept_of(col) in ("month", "quarter", "year", "date")


def _group_labels(values: pd.Series, col: str, lang: str) -> pd.Series:
    concept = _concept_of(col)
    if concept == "month" and pd.api.types.is_numeric_dtype(values):
        if lang == "vi":
            return values.map(lambda m: f"{T['vi']['month']} {int(m)}")
        return values.map(lambda m: MONTHS_EN[int(m) - 1] if 1 <= int(m) <= 12 else str(m))
    return values.astype(str)


def run_analysis(ctx: ToolContext, group_by: str, metric: str,
                 filters: dict | None = None) -> Analysis:
    df, meta = load_dataset(ctx.data_path)
    group_col = resolve_column(df, group_by)
    margin = _is_margin(metric)
    if margin:
        rev_col = resolve_column(df, "doanh thu", numeric=True)
        prof_col = resolve_column(df, "loi nhuan", numeric=True)
        metric_col, unit = "margin", "pct"
        metric_label = tr(ctx, "margin")
    else:
        metric_col = resolve_column(df, metric, numeric=True)
        unit = detect_unit(metric_col)
        metric_label = metric_col
    if group_col == metric_col:
        raise ToolError("group_by and metric must be different columns")

    df, applied = apply_filters(df, filters)
    year_col = meta.get("year_col")
    if not year_col and meta.get("date_col"):
        df = df.assign(__year=pd.to_datetime(df[meta["date_col"]]).dt.year)
        year_col = "__year"
    if group_col == meta.get("date_col"):
        df = df.assign(__period=pd.to_datetime(df[group_col]).dt.strftime("%Y-%m"))
        group_col_eff, year_split = "__period", False
    else:
        group_col_eff = group_col
        year_split = bool(year_col) and group_col != year_col and df[year_col].nunique() > 1

    def agg(frame: pd.DataFrame, by: list[str]) -> pd.Series:
        if margin:
            g = frame.groupby(by)[[rev_col, prof_col]].sum()
            return (g[prof_col] / g[rev_col].where(g[rev_col] != 0)).rename("value")
        return frame.groupby(by)[metric_col].sum()

    if year_split:
        years = sorted(int(y) for y in df[year_col].dropna().unique())
        s = agg(df, [group_col_eff, year_col])
        table = s.unstack(year_col).reindex(columns=years)
        last, prev = years[-1], years[-2]
        if margin:
            table["growth"] = table[last] - table[prev]  # percentage points
        else:
            table["growth"] = table[last] / table[prev].where(table[prev] != 0) - 1
        # Absolute change answers "which group contributed most to growth",
        # which the growth rate alone does not: a small group can double.
        table["change"] = table[last] - table[prev]
        latest = table[last]
    else:
        years = []
        s = agg(df, [group_col_eff])
        table = s.to_frame("value")
        table["growth"] = float("nan")
        table["change"] = float("nan")
        latest = table["value"]

    if margin:
        table["share"] = float("nan")
    else:
        total = latest.sum()
        table["share"] = latest / total if total else float("nan")

    time_like = _time_like(group_col) or group_col_eff == "__period"
    if time_like:
        table = table.sort_index()
    else:
        table = table.sort_values(years[-1] if years else "value", ascending=False)
    labels = _group_labels(pd.Series(table.index, index=table.index), group_col, ctx.lang)
    table.index = labels.values
    table.index.name = group_col

    return Analysis(group_col=group_col, metric_col=metric_col, metric_label=metric_label,
                    unit=unit, years=years, table=table, filters=applied,
                    rows_used=len(df), time_like=time_like, frame=df)


def _fmt(a: Analysis, v: float, lang: str, with_unit: bool = False) -> str:
    if a.unit == "pct":
        return fmt_pct(v, lang, signed=False)
    s = fmt_value(v, a.unit, lang)
    if with_unit and not a.unit and _concept_of(a.metric_col) == "units":
        s += " " + T.get(lang, T["en"])["units_word"]
    return s


def _fmt_change(a: Analysis, v: float, lang: str) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "—"
    s = _fmt(a, abs(v), lang, with_unit=True)
    return ("+" if v > 0 else "-" if v < 0 else "") + s


def metric_short(a: Analysis, lang: str) -> str:
    """'Doanh thu (triệu VND)' -> 'doanh thu', used after "% of total ..."."""
    if a.unit == "pct":
        return tr_lang(lang, "margin").lower()
    known = {"vi": {"revenue": "doanh thu", "profit": "lợi nhuận gộp", "units": "số lượng",
                    "cost": "giá vốn"},
             "en": {"revenue": "revenue", "profit": "gross profit", "units": "units",
                    "cost": "cost"}}
    concept = _concept_of(a.metric_col)
    word = known.get(lang, known["en"]).get(concept or "")
    return word or re.sub(r"\s*\(.*?\)\s*", " ", a.metric_label).strip().lower()


def tr_lang(lang: str, key: str) -> str:
    return T.get(lang, T["en"]).get(key, T["en"][key])


def _fmt_growth(a: Analysis, g: float, lang: str) -> str:
    if g is None or (isinstance(g, float) and math.isnan(g)):
        return "—"
    if a.unit == "pct":  # margin change, in percentage points
        pts = _group_digits(g * 100, 1, lang)
        return ("+" if g > 0 else "") + pts + " " + T.get(lang, T["en"])["pp"]
    return fmt_pct(g, lang)


def highlights(ctx: ToolContext, a: Analysis) -> list[str]:
    lang, t = ctx.lang, T.get(ctx.lang, T["en"])
    out: list[str] = []
    tab = a.table
    if tab.empty:
        return out
    col = a.years[-1] if a.years else "value"
    latest = tab[col].dropna()
    if latest.empty:
        return out
    top = latest.idxmax()
    share = tab.loc[top, "share"]
    tail = (f", {fmt_pct(share, lang, signed=False)} {t['of_total']} {metric_short(a, lang)}"
            if not math.isnan(share) else "")
    when = f" {a.years[-1]}" if a.years else ""
    out.append(f"{t['largest']}{when}: {top} ({_fmt(a, latest[top], lang, True)}{tail})")
    if len(latest) > 2:
        low = latest.idxmin()
        out.append(f"{t['smallest']}{when}: {low} ({_fmt(a, latest[low], lang, True)})")
    if a.years and len(a.years) > 1 and tab["growth"].notna().any():
        g = tab["growth"].dropna()
        # Word the extremes so they read correctly when everything shrank.
        if a.unit == "pct":
            best, worst = t["best_change"], t["worst_change"]
        else:
            best = t["fastest"] if g.max() >= 0 else t["least_decline"]
            worst = t["most_decline"] if g.min() < 0 else t["slowest"]
        out.append(f"{best}: {g.idxmax()} ({_fmt_growth(a, g.max(), lang)} "
                   f"{t['vs']} {a.years[-2]})")
        if len(g) > 1:
            out.append(f"{worst}: {g.idxmin()} ({_fmt_growth(a, g.min(), lang)} "
                       f"{t['vs']} {a.years[-2]})")
        if a.unit != "pct" and len(tab) > 1:
            ch = tab["change"].dropna()
            if len(ch) and ch.max() > 0:
                out.append(f"{t['biggest_gain']}: {ch.idxmax()} ({_fmt_change(a, ch.max(), lang)} "
                           f"{t['vs']} {a.years[-2]})")
            if len(ch) and ch.min() < 0:
                out.append(f"{t['biggest_loss']}: {ch.idxmin()} ({_fmt_change(a, ch.min(), lang)} "
                           f"{t['vs']} {a.years[-2]})")
        if a.unit != "pct":
            tot_last, tot_prev = tab[a.years[-1]].sum(), tab[a.years[-2]].sum()
            gr = tot_last / tot_prev - 1 if tot_prev else float("nan")
            out.append(f"{t['total']} {a.years[-1]}: {_fmt(a, tot_last, lang, True)} "
                       f"({_fmt_growth(a, gr, lang)} {t['vs']} {a.years[-2]}: "
                       f"{_fmt(a, tot_prev, lang, True)})")
    elif a.unit != "pct":
        out.append(f"{t['total']}: {_fmt(a, latest.sum(), lang, True)}")
    return out


def table_rows(ctx: ToolContext, a: Analysis, limit: int | None = None) -> list[dict]:
    t = T.get(ctx.lang, T["en"])
    rows = []
    for label, r in a.table.iterrows():
        row: dict[str, Any] = {a.group_col: label}
        if a.years:
            for y in a.years:
                row[str(y)] = _fmt(a, r[y], ctx.lang)
            if len(a.years) > 1:
                row[t["growth"]] = _fmt_growth(a, r["growth"], ctx.lang)
                if a.unit != "pct":
                    row[t["change"]] = _fmt_change(a, r["change"], ctx.lang)
        else:
            row[t["value"]] = _fmt(a, r["value"], ctx.lang)
        if not math.isnan(r["share"]):
            row[f"{t['share']} {a.years[-1] if a.years else ''}".strip()] = \
                fmt_pct(r["share"], ctx.lang, signed=False)
        rows.append(row)
        if limit and len(rows) >= limit:
            break
    return rows


def _ui_table(ctx: ToolContext, a: Analysis) -> dict:
    """Full table for the workshop UI: header plus formatted rows."""
    rows = table_rows(ctx, a)
    header = list(rows[0].keys()) if rows else [a.group_col]
    return {"columns": header, "rows": [[r.get(h, "") for h in header] for r in rows]}


# --------------------------------------------------------------------------
# Tool 1 — get_dataset_info
# --------------------------------------------------------------------------
def get_dataset_info(ctx: ToolContext) -> tuple[dict, dict]:
    df, meta = load_dataset(ctx.data_path)
    cols = []
    for c in df.columns:
        s = df[c]
        info: dict[str, Any] = {"name": c}
        if pd.api.types.is_datetime64_any_dtype(s):
            info["type"] = "date"
            info["range"] = f"{s.min():%Y-%m-%d} → {s.max():%Y-%m-%d}"
        elif pd.api.types.is_numeric_dtype(s):
            info["type"] = "number"
            if s.nunique() <= 12 and _concept_of(c) in ("year", "month", "quarter"):
                info["values"] = sorted(int(v) for v in s.dropna().unique())
            else:
                unit = detect_unit(c)
                info["total"] = fmt_value(float(s.sum()), unit, ctx.lang)
                info["min"] = fmt_value(float(s.min()), unit, ctx.lang)
                info["max"] = fmt_value(float(s.max()), unit, ctx.lang)
        else:
            info["type"] = "text"
            uniq = s.dropna().astype(str).unique()
            info["distinct"] = int(len(uniq))
            info["values" if len(uniq) <= 12 else "examples"] = list(uniq[:12])
        cols.append(info)
    time_range = None
    if meta.get("date_col"):
        d = pd.to_datetime(df[meta["date_col"]])
        time_range = f"{d.min():%Y-%m} → {d.max():%Y-%m}"
    years = sorted(int(y) for y in df[meta["year_col"]].dropna().unique()) if meta.get("year_col") else []
    numeric = [c for c in df.columns if pd.api.types.is_numeric_dtype(df[c]) and not _time_like(c)]
    categorical = [c for c in df.columns if not pd.api.types.is_numeric_dtype(df[c])
                   and not pd.api.types.is_datetime64_any_dtype(df[c])]
    has_margin = any(_concept_of(c) == "revenue" for c in numeric) and \
        any(_concept_of(c) == "profit" for c in numeric)
    model = {
        "file": meta["file"],
        "sheet_used": meta["sheet"],
        "sheets": meta["sheets"],
        "rows": int(len(df)),
        "time_range": time_range,
        "years": years,
        "columns": cols,
        "good_group_by": categorical + [c for c in df.columns if _concept_of(c) in ("month", "quarter")],
        "good_metrics": numeric + (["margin"] if has_margin else []),
    }
    ui = {"summary": f"{meta['file']} · {len(df):,} rows · {len(df.columns)} columns"
                     + (f" · {time_range}" if time_range else ""),
          "columns": [c["name"] for c in cols]}
    return model, ui


# --------------------------------------------------------------------------
# Tool 2 — analyze_sales
# --------------------------------------------------------------------------
def analyze_sales(ctx: ToolContext, group_by: str, metric: str,
                  filters: dict | None = None, top_n: int | None = None) -> tuple[dict, dict]:
    a = run_analysis(ctx, group_by, metric, filters)
    rows = table_rows(ctx, a, limit=min(top_n or MAX_MODEL_ROWS, MAX_MODEL_ROWS))
    hl = highlights(ctx, a)
    model = {
        "analysis_id": ctx.next_id("A"),
        "group_by": a.group_col,
        "metric": a.metric_label,
        "filters": a.filters or None,
        "years": a.years or None,
        "rows_used": a.rows_used,
        "table": rows,
        "highlights": hl,
        "note": "Numbers are already computed and formatted. Quote them exactly; do not recalculate.",
    }
    if len(a.table) > len(rows):
        model["table_truncated"] = f"{len(rows)} of {len(a.table)} groups shown"
    ui = {"table": _ui_table(ctx, a), "highlights": hl,
          "summary": f"{a.group_col} × {a.metric_label}"}
    return model, ui


# --------------------------------------------------------------------------
# Tool 3 — create_chart
# --------------------------------------------------------------------------
_mpl_lock = threading.Lock()
BRAND = "#0f8a63"
BRAND_DARK = "#0b6e4f"
PREV = "#9fb3c8"
PALETTE = ["#0f8a63", "#1d4ed8", "#b45309", "#7c3aed", "#db2777", "#0891b2", "#65a30d", "#dc2626"]


def _axis_scale(a: Analysis, lang: str) -> tuple[float, str]:
    t = T.get(lang, T["en"])
    if a.unit == "million_vnd":
        return 1000.0, t["billion"]
    if a.unit == "vnd":
        return 1e9, t["billion"]
    if a.unit == "pct":
        return 0.01, "%"
    return 1.0, ""


def draw_chart(ctx: ToolContext, a: Analysis, chart_type: str, title: str | None,
               path: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter

    lang = ctx.lang
    scale, unit_label = _axis_scale(a, lang)
    tab = a.table
    labels = [str(i) for i in tab.index]
    title = title or f"{a.metric_label} — {a.group_col}"

    with _mpl_lock:
        plt.rcParams.update({
            "font.family": "DejaVu Sans",  # bundled with matplotlib; covers Vietnamese
            "font.size": 11,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.edgecolor": "#cbd5e1",
            "axes.titleweight": "bold",
            "axes.titlesize": 15,
        })
        fig, ax = plt.subplots(figsize=(10, 5.6), dpi=130)
        fmt_axis = FuncFormatter(
            lambda v, _: _group_digits(v, 0 if abs(v) >= 10 or v == 0 else 1, lang))
        # Twelve "Tháng N" labels collide on a line chart; number the axis instead.
        month_axis = _concept_of(a.group_col) == "month" and len(labels) > 6
        tick_labels = ([str(i + 1) for i in range(len(labels))] if month_axis else labels)

        if chart_type == "pie":
            col = a.years[-1] if a.years else "value"
            vals = tab[col].clip(lower=0).fillna(0)
            colors = PALETTE * (len(vals) // len(PALETTE) + 1)
            wedges, _, autotexts = ax.pie(
                vals.values, labels=labels, colors=colors[: len(vals)], startangle=90,
                counterclock=False, autopct=lambda p: fmt_pct(p / 100, lang, signed=False),
                pctdistance=0.78, wedgeprops={"width": 0.45, "edgecolor": "white"},
                textprops={"fontsize": 10.5})
            for t_ in autotexts:
                t_.set_color("white")
                t_.set_fontweight("bold")
            ax.axis("equal")
            if a.years:
                title = f"{title} ({a.years[-1]})"
        elif chart_type == "line":
            x = range(len(labels))
            if a.years:
                for i, y in enumerate(a.years):
                    color = BRAND if y == a.years[-1] else PREV
                    ax.plot(x, tab[y].values / scale, marker="o", linewidth=2.6 if y == a.years[-1] else 2,
                            color=color, label=str(y))
            else:
                ax.plot(x, tab["value"].values / scale, marker="o", linewidth=2.6, color=BRAND)
            ax.set_xticks(list(x))
            ax.set_xticklabels(tick_labels, rotation=0 if len(labels) <= 12 else 45,
                               ha="center" if len(labels) <= 12 else "right")
            if month_axis:
                ax.set_xlabel(T.get(lang, T["en"])["month"])
            ax.yaxis.set_major_formatter(fmt_axis)
            ax.set_ylabel(unit_label)
            ax.grid(axis="y", color="#e2e8f0")
            if a.years:
                ax.legend(frameon=False)
        else:  # bar
            import numpy as np
            n = len(labels)
            x = np.arange(n)
            if a.years and len(a.years) > 1:
                shown = a.years[-2:]
                w = 0.38
                ax.bar(x - w / 2, tab[shown[0]].values / scale, w, color=PREV, label=str(shown[0]))
                bars = ax.bar(x + w / 2, tab[shown[1]].values / scale, w, color=BRAND, label=str(shown[1]))
                for rect, g in zip(bars, tab["growth"].values):
                    if g is None or (isinstance(g, float) and math.isnan(g)):
                        continue
                    ax.annotate(_fmt_growth(a, g, lang), (rect.get_x() + rect.get_width() / 2, rect.get_height()),
                                xytext=(0, 4), textcoords="offset points", ha="center", fontsize=9.5,
                                fontweight="bold", color=BRAND_DARK if g >= 0 else "#b91c1c")
                ax.legend(frameon=False)
            else:
                col = a.years[-1] if a.years else "value"
                bars = ax.bar(x, tab[col].values / scale, 0.6, color=BRAND)
                for rect, v in zip(bars, tab[col].values):
                    ax.annotate(_fmt(a, v, lang), (rect.get_x() + rect.get_width() / 2, rect.get_height()),
                                xytext=(0, 4), textcoords="offset points", ha="center", fontsize=9)
            ax.set_xticks(x)
            rotate = not month_axis and (n > 6 or max((len(s) for s in labels), default=0) > 14)
            ax.set_xticklabels(tick_labels, rotation=25 if rotate else 0, ha="right" if rotate else "center")
            if month_axis:
                ax.set_xlabel(T.get(lang, T["en"])["month"])
            ax.yaxis.set_major_formatter(fmt_axis)
            ax.set_ylabel(unit_label)
            ax.grid(axis="y", color="#e2e8f0")
            ax.set_axisbelow(True)

        ax.set_title(title, loc="left", pad=14)
        note = f"{T.get(lang, T['en'])['source']}: {Path(ctx.data_path).name}"
        if a.filters:
            note += " · " + ", ".join(f"{k} = {v}" for k, v in a.filters.items())
        fig.text(0.01, 0.01, note, fontsize=8.5, color="#6b7f95")
        fig.tight_layout(rect=(0, 0.03, 1, 1))
        path.parent.mkdir(parents=True, exist_ok=True)
        fig.savefig(path, facecolor="white")
        plt.close(fig)


def create_chart(ctx: ToolContext, group_by: str, metric: str, chart_type: str = "bar",
                 title: str | None = None, filters: dict | None = None) -> tuple[dict, dict]:
    if chart_type not in ("bar", "line", "pie"):
        chart_type = "bar"
    a = run_analysis(ctx, group_by, metric, filters)
    if chart_type == "pie" and a.unit == "pct":
        chart_type = "bar"  # a pie of percentages is meaningless
    chart_id = ctx.next_id("C")
    path = Path(ctx.out_dir) / f"chart-{chart_id}.png"
    draw_chart(ctx, a, chart_type, title, path)
    rec = {"id": chart_id, "path": str(path), "group_by": a.group_col, "metric": a.metric_col,
           "filters": a.filters, "title": title or ""}
    ctx.charts.append(rec)
    ctx.files.append({"kind": "chart", "name": path.name, "path": str(path)})
    model = {"chart_id": chart_id, "file": path.name, "chart_type": chart_type,
             "status": "created", "shows": f"{a.metric_label} by {a.group_col}"}
    ui = {"image": path.name, "title": title or model["shows"]}
    return model, ui


# --------------------------------------------------------------------------
# Tool 4 — export_excel_report
# --------------------------------------------------------------------------
def _safe_filename(text: str) -> str:
    s = norm(text).replace(" ", "-")[:60].strip("-")
    return s or "report"


def export_excel_report(ctx: ToolContext, title: str, group_by: str, metric: str,
                        insights: list[str] | None = None,
                        filters: dict | None = None) -> tuple[dict, dict]:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    t = T.get(ctx.lang, T["en"])
    a = run_analysis(ctx, group_by, metric, filters)
    insights = [str(i).strip() for i in (insights or []) if str(i).strip()][:8]
    report_id = ctx.next_id("R")

    head_fill = PatternFill("solid", fgColor="0B6E4F")
    soft_fill = PatternFill("solid", fgColor="E8F5F0")
    white_bold = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="CBD5E1")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()

    # -- Summary -----------------------------------------------------------
    ws = wb.active
    ws.title = t["summary"]
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=18, color="0D1B2A")
    ws["A2"] = t["made_by"]
    ws["A2"].font = Font(italic=True, size=10, color="6B7F95")
    meta_rows = [
        (t["dataset"], Path(ctx.data_path).name),
        (t["group_by"], a.group_col),
        (t["metric"], a.metric_label),
        (t["filters"], ", ".join(f"{k} = {v}" for k, v in a.filters.items()) or t["all_rows"]),
        (t["rows_used"], a.rows_used),
        (t["generated"], time.strftime("%Y-%m-%d %H:%M")),
    ]
    r = 4
    for k, v in meta_rows:
        ws.cell(r, 1, k).font = Font(bold=True, color="33475B")
        ws.cell(r, 2, v)
        r += 1
    r += 1
    ws.cell(r, 1, t["kpis"]).font = Font(bold=True, size=13, color="0B6E4F")
    r += 1
    for line in highlights(ctx, a):
        ws.cell(r, 1, "•")
        ws.cell(r, 2, line)
        r += 1
    r += 1
    ws.cell(r, 1, t["insights"]).font = Font(bold=True, size=13, color="0B6E4F")
    r += 1
    for line in insights or ["—"]:
        ws.cell(r, 1, "•")
        c = ws.cell(r, 2, line)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = max(18, 15 * (1 + len(line) // 95))
        r += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100

    # -- Analysis table + native Excel chart ---------------------------------
    wa = wb.create_sheet(t["analysis"])
    header = [a.group_col]
    value_cols = [str(y) for y in a.years] if a.years else [t["value"]]
    header += value_cols
    has_growth = len(a.years) > 1
    if has_growth:
        header.append(t["growth"])
    has_share = a.table["share"].notna().any()
    if has_share:
        header.append(t["share"])
    for j, h in enumerate(header, 1):
        c = wa.cell(1, j, h)
        c.font, c.fill, c.border = white_bold, head_fill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    scale, unit_label = (1.0, "")
    num_fmt = "#,##0"
    if a.unit == "million_vnd":
        scale, unit_label, num_fmt = 1000.0, t["billion"], "#,##0.0"
    elif a.unit == "pct":
        num_fmt = "0.0%"
    for i, (label, row) in enumerate(a.table.iterrows(), 2):
        wa.cell(i, 1, str(label)).border = box
        src = a.years if a.years else ["value"]
        for j, y in enumerate(src, 2):
            v = row[y]
            c = wa.cell(i, j, None if pd.isna(v) else float(v) / (scale if a.unit != "pct" else 1))
            c.number_format, c.border = num_fmt, box
        j = 2 + len(src)
        if has_growth:
            g = row["growth"]
            c = wa.cell(i, j, None if pd.isna(g) else float(g))
            c.number_format, c.border = ('+0.0%;-0.0%;0.0%' if a.unit != "pct" else '+0.0%;-0.0%'), box
            if not pd.isna(g):
                c.font = Font(bold=True, color="0B6E4F" if g >= 0 else "B91C1C")
            j += 1
        if has_share:
            s = row["share"]
            c = wa.cell(i, j, None if pd.isna(s) else float(s))
            c.number_format, c.border = "0.0%", box
        if i % 2 == 0:
            for jj in range(1, len(header) + 1):
                wa.cell(i, jj).fill = soft_fill
    last_row = 1 + len(a.table)
    if unit_label:
        wa.cell(last_row + 2, 1, f"({a.metric_label} — {unit_label})").font = Font(italic=True, color="6B7F95")
    for j in range(1, len(header) + 1):
        wa.column_dimensions[get_column_letter(j)].width = 26 if j == 1 else 16
    wa.freeze_panes = "B2"
    if len(a.table) >= 1:
        chart = LineChart() if a.time_like else BarChart()
        chart.title = title
        chart.y_axis.title = unit_label or a.metric_label
        chart.height, chart.width = 9, 18
        data = Reference(wa, min_col=2, max_col=1 + len(value_cols), min_row=1, max_row=last_row)
        cats = Reference(wa, min_col=1, min_row=2, max_row=last_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        wa.add_chart(chart, f"{get_column_letter(len(header) + 2)}2")

    # -- Chart image (matplotlib) -------------------------------------------
    match = next((c for c in reversed(ctx.charts)
                  if c["group_by"] == a.group_col and c["metric"] == a.metric_col
                  and c.get("filters") == a.filters), None)
    img_path = Path(match["path"]) if match else Path(ctx.out_dir) / f"chart-{report_id}.png"
    if not match:
        draw_chart(ctx, a, "line" if a.time_like else "bar", title, img_path)
    wc = wb.create_sheet(t["chart"])
    try:
        img = XLImage(str(img_path))
        img.width, img.height = 900, 504
        wc.add_image(img, "B2")
    except Exception as exc:  # noqa: BLE001 - Pillow missing: keep the report usable
        wc["B2"] = f"(chart image unavailable: {exc})"

    # -- Underlying rows -------------------------------------------------------
    wd = wb.create_sheet(t["data"])
    frame = a.frame.drop(columns=[c for c in a.frame.columns if c.startswith("__")])
    max_rows = 20000
    for j, h in enumerate(frame.columns, 1):
        c = wd.cell(1, j, h)
        c.font, c.fill = white_bold, head_fill
        wd.column_dimensions[get_column_letter(j)].width = max(10, min(30, len(str(h)) + 4))
    for i, rec in enumerate(frame.head(max_rows).itertuples(index=False), 2):
        for j, v in enumerate(rec, 1):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
                wd.cell(i, j, v).number_format = "yyyy-mm"
            elif pd.isna(v) if not isinstance(v, (list, dict)) else False:
                continue
            else:
                wd.cell(i, j, v.item() if hasattr(v, "item") else v)
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = wd.dimensions
    if len(frame) > max_rows:
        wd.cell(max_rows + 3, 1, t["truncated"].format(n=max_rows))

    name = f"{_safe_filename(title)}-{report_id}.xlsx"
    path = Path(ctx.out_dir) / name
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    ctx.files.append({"kind": "report", "name": name, "path": str(path)})
    model = {"report_id": report_id, "file": name,
             "sheets": [t["summary"], t["analysis"], t["chart"], t["data"]],
             "insights_written": len(insights), "status": "created"}
    ui = {"file": name, "sheets": model["sheets"], "title": title}
    return model, ui


# --------------------------------------------------------------------------
# Dispatcher
# --------------------------------------------------------------------------
def _coerce_args(args: Any) -> dict:
    if args is None or args == "":
        return {}
    if isinstance(args, str):
        try:
            args = json.loads(args)
        except json.JSONDecodeError as exc:
            raise ToolError(f"arguments are not valid JSON: {exc}") from exc
    if not isinstance(args, dict):
        raise ToolError("arguments must be a JSON object")
    # Small models sometimes send filters as a JSON string.
    if isinstance(args.get("filters"), str):
        try:
            args["filters"] = json.loads(args["filters"]) if args["filters"].strip() else None
        except json.JSONDecodeError:
            args["filters"] = None
    if isinstance(args.get("insights"), str):
        args["insights"] = [s.strip(" -•") for s in re.split(r"\n+", args["insights"]) if s.strip()]
    return args


def run_tool(name: str, args: Any, ctx: ToolContext) -> tuple[dict, dict, bool]:
    """Execute one tool call. Returns (payload for the model, payload for the UI, ok)."""
    try:
        a = _coerce_args(args)
        if name == "get_dataset_info":
            model, ui = get_dataset_info(ctx)
        elif name == "analyze_sales":
            model, ui = analyze_sales(ctx, a.get("group_by", ""), a.get("metric", ""),
                                      a.get("filters"), a.get("top_n"))
        elif name == "create_chart":
            model, ui = create_chart(ctx, a.get("group_by", ""), a.get("metric", ""),
                                     a.get("chart_type", "bar"), a.get("title"), a.get("filters"))
        elif name == "export_excel_report":
            model, ui = export_excel_report(ctx, a.get("title") or "Báo cáo phân tích doanh số",
                                            a.get("group_by", ""), a.get("metric", ""),
                                            a.get("insights"), a.get("filters"))
        else:
            raise ToolError(f"unknown tool '{name}'", available_tools=TOOL_NAMES)
        return model, ui, True
    except ToolError as exc:
        err = {"error": str(exc), **exc.extra}
        return err, err, False
    except TypeError as exc:
        err = {"error": f"bad arguments for {name}: {exc}"}
        return err, err, False


# --------------------------------------------------------------------------
# CLI — used by the OpenClaw skill inside the sandbox
# --------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Sales analyst tools (pandas / matplotlib / openpyxl)")
    ap.add_argument("tool", choices=TOOL_NAMES + ["specs"])
    ap.add_argument("--args", default="{}", help="tool arguments as a JSON object")
    ap.add_argument("--data", default=str(DEFAULT_DATA), help="the .xlsx or .csv file")
    ap.add_argument("--out", default="outputs", help="where charts and reports are written")
    ap.add_argument("--lang", default="vi", choices=["vi", "en"])
    ns = ap.parse_args(argv)
    if ns.tool == "specs":
        print(json.dumps(TOOL_SPECS, ensure_ascii=False, indent=2))
        return 0
    out = Path(ns.out).resolve()
    ctx = ToolContext(data_path=Path(ns.data).resolve(), out_dir=out, lang=ns.lang)
    # Number ids by what is already in the output folder, so repeated CLI calls
    # do not overwrite earlier charts.
    if out.exists():
        ctx.counter = len(list(out.glob("chart-*.png"))) + len(list(out.glob("*.xlsx")))
    model, _, ok = run_tool(ns.tool, ns.args, ctx)
    for f in ctx.files:
        model.setdefault("files", []).append(f["path"])
    print(json.dumps(model, ensure_ascii=False, indent=2, default=str))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
